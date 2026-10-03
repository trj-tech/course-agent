"""课程资料：上传、解析、切片向量化（RAG 入库）、列表。"""
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.core.security import get_current_user
from app.database import get_db
from app.models.document import CourseDocument
from app.models.user import User
from app.rag.indexer import index_document

router = APIRouter(prefix="/documents", tags=["documents"])

UPLOAD_DIR = Path(__file__).resolve().parent.parent.parent / "data" / "uploads"
ALLOWED = {".pdf", ".docx", ".txt", ".md"}


def extract_text(path: Path, file_type: str) -> str:
    """按文件类型抽取纯文本。"""
    if file_type == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if file_type == "docx":
        import docx

        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    return path.read_text(encoding="utf-8", errors="ignore")


@router.get("")
def list_documents(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    docs = db.query(CourseDocument).filter(CourseDocument.user_id == current_user.id).all()
    return [
        {
            "id": d.id,
            "filename": d.filename,
            "file_type": d.file_type,
            "status": d.status,
            "created_at": str(d.created_at),
        }
        for d in docs
    ]


@router.post("/upload")
def upload_document(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED:
        raise HTTPException(status_code=400, detail=f"仅支持 {ALLOWED} 格式")

    user_dir = UPLOAD_DIR / str(current_user.id)
    user_dir.mkdir(parents=True, exist_ok=True)
    counter = len(list(user_dir.glob("*")))
    dest = user_dir / f"{current_user.id}_{counter}_{file.filename}"
    dest.write_bytes(file.file.read())

    doc = CourseDocument(
        user_id=current_user.id,
        filename=file.filename,
        file_path=str(dest),
        file_type=suffix.lstrip("."),
        status="pending",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        text = extract_text(dest, doc.file_type)
        index_document(db, doc, text)
    except Exception as e:  # noqa: BLE001
        doc.status = "failed"
        db.commit()
        raise HTTPException(status_code=500, detail=f"文档处理失败：{e}")

    return {"id": doc.id, "filename": doc.filename, "status": doc.status}
